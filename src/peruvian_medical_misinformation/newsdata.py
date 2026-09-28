"""Batch reproducible de candidatas de NewsData.io.

La API solo descubre candidatas por URL y título. No decide si una noticia es
médica ni asigna una etiqueta de veracidad.
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import time
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping
from urllib.parse import parse_qsl, urlencode, urljoin, urlsplit, urlunsplit

import httpx
import yaml
from bs4 import BeautifulSoup


CANDIDATE_FIELDS = [
    "run_id",
    "seeded_from_run_id",
    "record_id",
    "source_dataset",
    "source_id",
    "source_name",
    "source_domain",
    "url",
    "canonical_url",
    "title",
    "description",
    "author",
    "published_at",
    "language",
    "country",
    "category",
    "newsdata_article_id",
    "api_query",
    "retrieved_at",
]

RUN_REGISTRY_FIELDS = [
    "run_number",
    "run_id",
    "executed_at",
    "status",
    "expected_total",
    "collected_total",
    "shortfall_total",
    "total_requests",
    "rate_limited",
    "request_budget_exhausted",
    "seeded_total",
    "candidates_path",
    "summary_path",
    "resumed_from",
]

TRACKING_PARAMETERS = {"fbclid", "gclid", "mc_cid", "mc_eid"}
ARTICLE_URL_PATTERN = re.compile(r"_\d{8}/?$")


class NewsDataConfigurationError(ValueError):
    """La configuración no define un batch ejecutable."""


class NewsDataRequestError(RuntimeError):
    """NewsData devolvió un error que debe quedar registrado."""


class NewsDataRateLimitError(NewsDataRequestError):
    """La API agotó su ventana de créditos y la corrida debe detenerse."""


def load_yaml(path: str | Path) -> dict[str, Any]:
    with Path(path).open(encoding="utf-8") as handle:
        payload = yaml.safe_load(handle) or {}
    if not isinstance(payload, dict):
        raise NewsDataConfigurationError(f"{path} debe contener un objeto YAML")
    return payload


def load_local_env(path: str | Path = ".env") -> None:
    """Carga pares simples KEY=VALUE sin sobrescribir variables ya definidas."""

    env_file = Path(path)
    if not env_file.exists():
        return
    for raw_line in env_file.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


def api_key_from_environment() -> str:
    value = os.environ.get("NEWSDATA_API_KEY", "").strip()
    if not value:
        raise NewsDataConfigurationError(
            "Falta NEWSDATA_API_KEY. Copia .env.example como .env y agrega la clave localmente."
        )
    return value


def canonicalize_url(url: str) -> str:
    parsed = urlsplit(url.strip())
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.hostname:
        raise ValueError("URL inválida")
    query = [
        (key, value)
        for key, value in parse_qsl(parsed.query, keep_blank_values=True)
        if key.lower() not in TRACKING_PARAMETERS and not key.lower().startswith("utm_")
    ]
    hostname = parsed.hostname.lower()
    netloc = hostname if parsed.port is None else f"{hostname}:{parsed.port}"
    return urlunsplit((parsed.scheme.lower(), netloc, parsed.path or "/", urlencode(query), ""))


def record_id(canonical_url: str) -> str:
    return hashlib.sha256(canonical_url.encode("utf-8")).hexdigest()[:16]


def has_source_domain(url: str, domain: str) -> bool:
    hostname = (urlsplit(url).hostname or "").lower().rstrip(".")
    expected = domain.lower().lstrip(".").rstrip(".")
    return hostname == expected or hostname.endswith(f".{expected}")


def normalize_article(
    item: Mapping[str, Any],
    *,
    source_id: str,
    source: Mapping[str, Any],
    run_id: str,
    query: str,
    retrieved_at: str,
) -> dict[str, str] | None:
    raw_url = str(item.get("link") or item.get("url") or "").strip()
    if not raw_url:
        return None
    try:
        normalized_url = canonicalize_url(raw_url)
    except ValueError:
        return None
    domain = str(source["domain"])
    if not has_source_domain(normalized_url, domain):
        return None
    country = item.get("country") or []
    category = item.get("category") or []
    creator = item.get("creator") or item.get("author") or []
    return {
        "run_id": run_id,
        "seeded_from_run_id": "",
        "record_id": record_id(normalized_url),
        "source_dataset": str(source.get("newsdata_source_dataset") or source["source_dataset"]),
        "source_id": source_id,
        "source_name": str(source["name"]),
        "source_domain": domain,
        "url": raw_url,
        "canonical_url": normalized_url,
        "title": str(item.get("title") or "").strip(),
        "description": str(item.get("description") or "").strip(),
        "author": ", ".join(map(str, creator)) if isinstance(creator, list) else str(creator),
        "published_at": str(item.get("pubDate") or item.get("pub_date") or "").strip(),
        "language": str(item.get("language") or "").strip(),
        "country": ",".join(map(str, country)) if isinstance(country, list) else str(country),
        "category": ",".join(map(str, category)) if isinstance(category, list) else str(category),
        "newsdata_article_id": str(item.get("article_id") or "").strip(),
        "api_query": query,
        "retrieved_at": retrieved_at,
    }


def normalize_seeded_article(
    item: Mapping[str, Any],
    *,
    source_id: str,
    source: Mapping[str, Any],
    run_id: str,
) -> dict[str, str] | None:
    """Conserva una candidata previa al reanudar una corrida incompleta."""

    raw_url = str(item.get("url") or "").strip()
    if not raw_url:
        return None
    try:
        normalized_url = canonicalize_url(raw_url)
    except ValueError:
        return None
    domain = str(source["domain"])
    if not has_source_domain(normalized_url, domain):
        return None
    return {
        "run_id": run_id,
        "seeded_from_run_id": str(item.get("run_id") or "").strip(),
        "record_id": record_id(normalized_url),
        "source_dataset": str(source["source_dataset"]),
        "source_id": source_id,
        "source_name": str(source["name"]),
        "source_domain": domain,
        "url": raw_url,
        "canonical_url": normalized_url,
        "title": str(item.get("title") or "").strip(),
        "description": str(item.get("description") or "").strip(),
        "author": str(item.get("author") or "").strip(),
        "published_at": str(item.get("published_at") or "").strip(),
        "language": str(item.get("language") or "").strip(),
        "country": str(item.get("country") or "").strip(),
        "category": str(item.get("category") or "").strip(),
        "newsdata_article_id": str(item.get("newsdata_article_id") or "").strip(),
        "api_query": str(item.get("api_query") or "").strip(),
        "retrieved_at": str(item.get("retrieved_at") or "").strip(),
    }


def archive_title(anchor: Any) -> str:
    """Obtiene el título visible o el alt de la imagen dentro de un enlace."""

    text = anchor.get_text(" ", strip=True)
    if text:
        return text
    image = anchor.select_one("img[alt]")
    return str(image.get("alt") or "").strip() if image else ""


def discovery_text(value: str) -> str:
    """Normaliza mínimamente texto para el filtro de descubrimiento de URLs."""

    return "".join(
        character for character in unicodedata.normalize("NFKD", value.casefold()) if not unicodedata.combining(character)
    )


def normalize_archive_article(
    *,
    raw_url: str,
    title: str,
    source_id: str,
    source: Mapping[str, Any],
    run_id: str,
    archive_url: str,
    retrieved_at: str,
) -> dict[str, str] | None:
    """Convierte un enlace de archivo público en una candidata trazable."""

    try:
        normalized_url = canonicalize_url(raw_url)
    except ValueError:
        return None
    domain = str(source["domain"])
    path = urlsplit(normalized_url).path
    prefixes = [str(prefix) for prefix in source.get("article_url_prefixes", []) if str(prefix).strip()]
    is_article_url = (
        True
        if source.get("allow_selected_article_urls", False)
        else any(path.startswith(prefix) for prefix in prefixes)
        if prefixes
        else bool(ARTICLE_URL_PATTERN.search(path))
    )
    if not has_source_domain(normalized_url, domain) or not is_article_url:
        return None
    clean_title = title.strip()
    if not clean_title:
        return None
    keywords = [discovery_text(str(keyword)) for keyword in source.get("candidate_keywords", []) if str(keyword).strip()]
    if keywords and not any(keyword in discovery_text(f"{clean_title} {normalized_url}") for keyword in keywords):
        return None
    return {
        "run_id": run_id,
        "seeded_from_run_id": "",
        "record_id": record_id(normalized_url),
        "source_dataset": str(source["source_dataset"]),
        "source_id": source_id,
        "source_name": str(source["name"]),
        "source_domain": domain,
        "url": raw_url,
        "canonical_url": normalized_url,
        "title": clean_title,
        "description": "",
        "author": "",
        "published_at": "",
        "language": "es",
        "country": "pe",
        "category": "health_archive",
        "newsdata_article_id": "",
        "api_query": f"archive:{archive_url}",
        "retrieved_at": retrieved_at,
    }


def archive_candidates_from_html(
    html: str,
    *,
    archive_url: str,
    source_id: str,
    source: Mapping[str, Any],
    run_id: str,
    retrieved_at: str,
) -> list[dict[str, str]]:
    """Extrae enlaces de artículos, no menús ni rutas de archivo, de una página pública."""

    soup = BeautifulSoup(html, "html.parser")
    candidates: dict[str, dict[str, str]] = {}
    # Cada fuente declara el selector de sus tarjetas editoriales. Latina usa
    # main-card; Perú21 usa nodos Drupal dentro de view-content.
    selector = str(source.get("article_link_selector") or "section#principal figure.main-card a[href]")
    anchors = soup.select(selector)
    if not anchors:
        anchors = soup.select("a[href]")
    for anchor in anchors:
        raw_url = urljoin(archive_url, str(anchor["href"]))
        row = normalize_archive_article(
            raw_url=raw_url,
            title=archive_title(anchor),
            source_id=source_id,
            source=source,
            run_id=run_id,
            archive_url=archive_url,
            retrieved_at=retrieved_at,
        )
        if row is not None:
            existing = candidates.get(row["canonical_url"])
            # Una tarjeta puede enlazar la imagen y el titular. Se prefiere el
            # texto más informativo, para no conservar valores como "Imagen".
            if existing is None or len(row["title"]) > len(existing["title"]):
                candidates[row["canonical_url"]] = row
    return list(candidates.values())


def validate_batch_config(config: Mapping[str, Any]) -> tuple[Mapping[str, Any], Mapping[str, Mapping[str, Any]]]:
    batch = config.get("batch")
    sources = config.get("sources")
    if not isinstance(batch, Mapping) or not isinstance(sources, Mapping):
        raise NewsDataConfigurationError("Se requieren los mapas 'batch' y 'sources'")
    required = ("endpoint", "output_dir", "report_dir", "per_source_limit", "expected_total", "health_queries")
    missing = [name for name in required if not batch.get(name)]
    if missing:
        raise NewsDataConfigurationError("Faltan campos de batch: " + ", ".join(missing))
    valid_sources: dict[str, Mapping[str, Any]] = {}
    for source_id, source in sources.items():
        if not isinstance(source, Mapping) or not all(source.get(key) for key in ("name", "source_dataset", "domain")):
            raise NewsDataConfigurationError(f"Fuente inválida: {source_id}")
        discovery = str(source.get("discovery", "newsdata"))
        if discovery not in {"newsdata", "archive"}:
            raise NewsDataConfigurationError(f"Tipo de descubrimiento inválido para {source_id}: {discovery}")
        if discovery == "archive" and not source.get("archive_urls"):
            raise NewsDataConfigurationError(f"Faltan archive_urls para {source_id}")
        valid_sources[str(source_id)] = source
    if not valid_sources:
        raise NewsDataConfigurationError("No hay fuentes configuradas")
    expected = int(batch["expected_total"])
    per_source = int(batch["per_source_limit"])
    if expected != per_source * len(valid_sources):
        raise NewsDataConfigurationError(
            "expected_total debe ser per_source_limit multiplicado por el número de fuentes"
        )
    return batch, valid_sources


RequestJSON = Callable[[str, Mapping[str, str]], Mapping[str, Any]]
RequestHTML = Callable[[str], str]


def httpx_json_request(timeout_seconds: float) -> RequestJSON:
    def request(endpoint: str, params: Mapping[str, str]) -> Mapping[str, Any]:
        response = httpx.get(endpoint, params=params, timeout=timeout_seconds, follow_redirects=True)
        try:
            payload = response.json()
        except ValueError as error:
            raise NewsDataRequestError(f"Respuesta no JSON ({response.status_code})") from error
        if not isinstance(payload, Mapping):
            raise NewsDataRequestError(f"Respuesta JSON no válida ({response.status_code})")
        if response.status_code == 429:
            details = payload.get("results")
            message = details.get("message") if isinstance(details, Mapping) else ""
            raise NewsDataRateLimitError(message or "NewsData limitó temporalmente la cuenta")
        if response.status_code >= 400 or str(payload.get("status", "success")).lower() == "error":
            details = payload.get("results")
            message = details.get("message") if isinstance(details, Mapping) else ""
            raise NewsDataRequestError(f"NewsData respondió HTTP {response.status_code}: {message or payload}")
        return payload

    return request


def httpx_html_request(timeout_seconds: float) -> RequestHTML:
    """Lee HTML público para descubrir enlaces; no descarga cuerpos de artículos."""

    headers = {
        "User-Agent": "peruvian-medical-misinformation-research/0.1 "
        "(academic corpus discovery; contact: research@example.invalid)"
    }

    def request(url: str) -> str:
        response = httpx.get(url, headers=headers, timeout=timeout_seconds, follow_redirects=True)
        if response.status_code >= 400:
            raise NewsDataRequestError(f"Archivo público respondió HTTP {response.status_code}: {url}")
        return response.text

    return request


def collect_candidates(
    config: Mapping[str, Any],
    *,
    api_key: str,
    run_id: str,
    request_json: RequestJSON,
    request_html: RequestHTML | None = None,
    seeded_rows: Iterable[Mapping[str, Any]] = (),
    excluded_record_ids: Iterable[str] = (),
    sleep: Callable[[float], None] = time.sleep,
    now: Callable[[], datetime] = lambda: datetime.now(timezone.utc),
) -> tuple[list[dict[str, str]], dict[str, Any]]:
    """Recolecta candidatas nuevas y prueba otras consultas ante duplicados."""

    batch, sources = validate_batch_config(config)
    limit = int(batch["per_source_limit"])
    results_per_request = max(limit, int(batch.get("results_per_request", limit)))
    delay = float(batch.get("request_delay_seconds", 0))
    max_requests = int(batch.get("max_requests_per_source", 1))
    max_total_requests = int(batch.get("max_requests_total", 30))
    if max_total_requests < 1:
        raise NewsDataConfigurationError("max_requests_total debe ser al menos 1")
    queries = [str(value).strip() for value in batch["health_queries"] if str(value).strip()]
    excluded_ids = {str(value).strip() for value in excluded_record_ids if str(value).strip()}
    seeded_by_source: dict[str, dict[str, dict[str, str]]] = {source_id: {} for source_id in sources}
    for item in seeded_rows:
        source_id = str(item.get("source_id") or "").strip()
        source = sources.get(source_id)
        if source is None:
            continue
        row = normalize_seeded_article(item, source_id=source_id, source=source, run_id=run_id)
        if row is not None and row["record_id"] not in excluded_ids:
            seeded_by_source[source_id].setdefault(row["canonical_url"], row)

    candidates: list[dict[str, str]] = []
    errors: list[dict[str, str]] = []
    per_source: dict[str, dict[str, Any]] = {}
    total_requests = 0
    archive_requests_total = 0
    rate_limited = False
    request_budget_exhausted = False
    duplicates_skipped_ids: set[str] = set()

    for source_id, source in sources.items():
        selected = dict(seeded_by_source[source_id])
        requests_made = 0
        source_duplicate_ids: set[str] = set()
        source_status = "completed"
        discovery = str(source.get("discovery", "newsdata"))
        if discovery == "archive":
            if request_html is None:
                raise NewsDataConfigurationError("Se requiere request_html para una fuente de archivo público")
            for archive_url in source.get("archive_urls", []):
                if len(selected) >= limit:
                    break
                requests_made += 1
                archive_requests_total += 1
                try:
                    html = request_html(str(archive_url))
                except (httpx.HTTPError, NewsDataRequestError) as error:
                    errors.append({"source_id": source_id, "query": str(archive_url), "error": str(error)})
                    continue
                retrieved_at = now().isoformat()
                for row in archive_candidates_from_html(
                    html,
                    archive_url=str(archive_url),
                    source_id=source_id,
                    source=source,
                    run_id=run_id,
                    retrieved_at=retrieved_at,
                ):
                    if row["record_id"] in excluded_ids:
                        source_duplicate_ids.add(row["record_id"])
                        duplicates_skipped_ids.add(row["record_id"])
                        continue
                    selected.setdefault(row["canonical_url"], row)
                    if len(selected) >= limit:
                        break
                if delay and len(selected) < limit:
                    sleep(delay)
            rows = list(selected.values())[:limit]
            if len(selected) >= limit or not source.get("newsdata_fallback", False):
                candidates.extend(rows)
                per_source[source_id] = {
                    "requested": limit,
                    "collected": len(rows),
                    "shortfall": max(0, limit - len(rows)),
                    "requests_made": requests_made,
                    "seeded": len(seeded_by_source[source_id]),
                    "duplicates_skipped": len(source_duplicate_ids),
                    "status": source_status,
                }
                continue
        for query in queries:
            if len(selected) >= limit or requests_made >= max_requests:
                break
            if total_requests >= max_total_requests:
                request_budget_exhausted = True
                source_status = "request_budget_exhausted"
                break
            params = {
                "apikey": api_key,
                "domainurl": str(source.get("newsdata_domain") or source["domain"]),
                "language": str(batch.get("language", "es")),
                "country": str(batch.get("country", "pe")),
                "category": str(batch.get("category", "health")),
                "q": query,
                "size": str(results_per_request),
            }
            requests_made += 1
            total_requests += 1
            try:
                payload = request_json(str(batch["endpoint"]), params)
            except NewsDataRateLimitError as error:
                errors.append({"source_id": source_id, "query": query, "error": str(error)})
                rate_limited = True
                source_status = "rate_limited"
                break
            except (httpx.HTTPError, NewsDataRequestError) as error:
                errors.append({"source_id": source_id, "query": query, "error": str(error)})
                continue
            retrieved_at = now().isoformat()
            for item in payload.get("results", []):
                if not isinstance(item, Mapping):
                    continue
                row = normalize_article(
                    item,
                    source_id=source_id,
                    source=source,
                    run_id=run_id,
                    query=query,
                    retrieved_at=retrieved_at,
                )
                if row is not None:
                    if row["record_id"] in excluded_ids:
                        source_duplicate_ids.add(row["record_id"])
                        duplicates_skipped_ids.add(row["record_id"])
                        continue
                    selected.setdefault(row["canonical_url"], row)
                    if len(selected) >= limit:
                        break
            if delay and len(selected) < limit:
                sleep(delay)
        rows = list(selected.values())[:limit]
        candidates.extend(rows)
        per_source[source_id] = {
            "requested": limit,
            "collected": len(rows),
            "shortfall": max(0, limit - len(rows)),
            "requests_made": requests_made,
            "seeded": len(seeded_by_source[source_id]),
            "duplicates_skipped": len(source_duplicate_ids),
            "status": source_status,
        }
        if rate_limited or request_budget_exhausted:
            break

    for source_id in sources:
        if source_id not in per_source:
            per_source[source_id] = {
                "requested": limit,
                "collected": 0,
                "shortfall": limit,
                "requests_made": 0,
                "seeded": 0,
                "duplicates_skipped": 0,
                "status": "not_requested_after_rate_limit" if rate_limited else "not_requested_after_budget",
            }

    candidates.sort(key=lambda row: (row["source_id"], row["published_at"], row["canonical_url"]))
    summary: dict[str, Any] = {
        "run_id": run_id,
        "expected_total": int(batch["expected_total"]),
        "collected_total": len(candidates),
        "shortfall_total": max(0, int(batch["expected_total"]) - len(candidates)),
        "per_source": per_source,
        "errors": errors,
        "seeded_total": sum(len(rows) for rows in seeded_by_source.values()),
        "excluded_record_ids_total": len(excluded_ids),
        "duplicates_skipped_total": len(duplicates_skipped_ids),
        "total_requests": total_requests,
        "archive_requests_total": archive_requests_total,
        "rate_limited": rate_limited,
        "request_budget_exhausted": request_budget_exhausted,
    }
    return candidates, summary


def write_candidates(path: str | Path, rows: Iterable[Mapping[str, str]]) -> None:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=CANDIDATE_FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows({field: row.get(field, "") for field in CANDIDATE_FIELDS} for row in rows)


def read_candidates(path: str | Path) -> list[dict[str, str]]:
    """Lee una salida previa de candidatas para completar solo sus faltantes."""

    source = Path(path)
    with source.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        fieldnames = set(reader.fieldnames or [])
        missing = {"source_id", "url"} - fieldnames
        if missing:
            raise NewsDataConfigurationError(
                "El CSV para reanudar no contiene: " + ", ".join(sorted(missing))
            )
        return [dict(row) for row in reader]


def read_record_ids(path: str | Path) -> set[str]:
    """Lee record_id desde CSV o desde una hoja de cálculo de Excel."""

    source = Path(path)
    if source.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(source, read_only=True, data_only=True)
        worksheets = list(workbook.worksheets)
        if "Raw" in workbook.sheetnames:
            raw_sheet = workbook["Raw"]
            worksheets = [raw_sheet, *(sheet for sheet in worksheets if sheet.title != "Raw")]
        try:
            for worksheet in worksheets:
                rows = worksheet.iter_rows(values_only=True)
                header = next(rows, None)
                fieldnames = [str(value or "").strip() for value in header or ()]
                if "record_id" not in fieldnames:
                    continue
                record_id_index = fieldnames.index("record_id")
                return {
                    str(row[record_id_index] or "").strip()
                    for row in rows
                    if len(row) > record_id_index and str(row[record_id_index] or "").strip()
                }
        finally:
            workbook.close()
        raise NewsDataConfigurationError(
            "El Excel de registros existentes debe contener una columna record_id"
        )

    if source.suffix.lower() != ".csv":
        raise NewsDataConfigurationError(
            "El archivo de registros existentes debe ser .xlsx, .xlsm o .csv"
        )
    with source.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        fieldnames = set(reader.fieldnames or [])
        if "record_id" not in fieldnames:
            raise NewsDataConfigurationError(
                "El archivo de registros existentes debe contener la columna record_id"
            )
        return {
            str(row.get("record_id") or "").strip()
            for row in reader
            if str(row.get("record_id") or "").strip()
        }


def write_summary(path: str | Path, summary: Mapping[str, Any]) -> None:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_run_registry(path: str | Path) -> list[dict[str, str]]:
    """Lee el historial local de corridas; un registro ausente equivale a cero corridas."""

    registry = Path(path)
    if not registry.exists():
        return []
    with registry.open(newline="", encoding="utf-8") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def next_run_number(rows: Iterable[Mapping[str, str]]) -> int:
    numbers = []
    for row in rows:
        try:
            numbers.append(int(str(row.get("run_number") or "")))
        except ValueError:
            continue
    return max(numbers, default=0) + 1


def run_status(summary: Mapping[str, Any]) -> str:
    if bool(summary.get("rate_limited")):
        return "partial_rate_limited"
    if int(summary.get("collected_total", 0)) == int(summary.get("expected_total", 0)):
        return "complete"
    if bool(summary.get("request_budget_exhausted")):
        return "partial_request_budget"
    return "partial_no_results"


def append_run_registry(path: str | Path, row: Mapping[str, object]) -> None:
    """Agrega una corrida ya terminada al contador local sin sobrescribir el historial."""

    registry = Path(path)
    registry.parent.mkdir(parents=True, exist_ok=True)
    write_header = not registry.exists() or registry.stat().st_size == 0
    with registry.open("a", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=RUN_REGISTRY_FIELDS, lineterminator="\n")
        if write_header:
            writer.writeheader()
        writer.writerow({field: row.get(field, "") for field in RUN_REGISTRY_FIELDS})
